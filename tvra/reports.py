"""Reviewable XLSX and DOCX exports with explicit evidence and data limits.

Export functions never choose a directory or input file on the caller's behalf.
Existing files are protected unless allow_overwrite is explicitly supplied;
listed source_paths/input_paths are always protected.
"""
from __future__ import annotations

import json
import math
import re
from collections import Counter
from dataclasses import asdict, is_dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from tvra import RULE_VERSION


_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_LEVEL_NAMES = {1: "一级线索（优先复核）", 2: "二级线索", 3: "三级线索"}
_CAUTION = "本报告为基于已导入过车观测和可解释规则的辅助研判材料。风险线索不等于违法事实，应结合原始记录和人工核查作出判断。"
_LIMITS = [
    "卡口记录是离散过车观测，不能单独证明车辆实际停车、装卸、运输物品或违法行为。",
    "未被观测的路段、时间和行为无法还原；未命中规则不代表没有异常。",
    "疑似停留需要相邻观测的卡口或坐标支持，行政区域相同及时间跨度本身不能证明停留。",
    "缺少卡口、坐标、地点标签、出入口类型或所属单位时，相关规则可能未评估；未评估与未发现异常应区分。",
    "分析历史只记录导入或重新分析操作，重复分析不代表新增行程或新增异常。",
    "规则阈值和风险等级是筛查参考，需用真实核查结果持续评估，不是经过校准的违法概率。",
    "同车时间重叠的事件共享事件组融合分；事件分值不可再次相加。多条规则可引用相同观测记录。",
]


def _mapping(value):
    if value is None:
        return {}
    if is_dataclass(value):
        return asdict(value)
    if isinstance(value, dict):
        return value
    return dict(value)


def _text(value, default="未提供"):
    if value is None or value == "":
        return default
    if isinstance(value, (list, tuple, set)):
        value = "；".join(_text(item, "") for item in value)
    elif isinstance(value, dict):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    return _INVALID_XML.sub("", str(value))


def _number(value, default=0):
    try:
        result = float(value)
        return result if math.isfinite(result) else default
    except (TypeError, ValueError, OverflowError):
        return default


def _level(value):
    try:
        result = int(value)
    except (ValueError, TypeError):
        return None
    return result if result in (0, 1, 2, 3) else None


def _status(vehicle):
    if vehicle.get("assessment_status"):
        return _text(vehicle["assessment_status"])
    if _number(vehicle.get("record_count")) <= 0:
        return "未评估：无可用过车记录"
    if _number(vehicle.get("event_count")) > 0 or _level(vehicle.get("risk_level")) in (1, 2, 3):
        return "发现疑似异常线索，待人工核查"
    if _level(vehicle.get("risk_level")) == 0:
        return "已运行规则未发现异常；字段不足的规则可能未评估"
    return "未评估：未提供分析结果"


def _level_text(vehicle):
    level = _level(vehicle.get("risk_level", vehicle.get("level")))
    if _number(vehicle.get("record_count", 1)) <= 0:
        return "未评估"
    return _LEVEL_NAMES.get(level, "未发现规则命中" if level == 0 else "未评估")


def _output_path(path, suffix, metadata):
    target = Path(path).expanduser()
    if target.suffix.lower() != suffix:
        raise ValueError("输出文件扩展名应为 " + suffix)
    target = target.resolve()
    sources = []
    for key in ("source_paths", "input_paths"):
        values = metadata.get(key, [])
        if isinstance(values, (str, Path)):
            values = [values]
        sources.extend(values or [])
    for source in sources:
        source_path = Path(source).expanduser().resolve()
        if target == source_path or (
            target.exists() and source_path.exists() and target.samefile(source_path)
        ):
            raise ValueError("不能将报告写入导入源文件，请选择新的输出文件名")
    if target.exists() and not metadata.get("allow_overwrite", False):
        raise FileExistsError("输出文件已存在，请选择新文件名：" + str(target))
    if not target.parent.is_dir():
        raise FileNotFoundError("输出目录不存在：" + str(target.parent))
    return target


def _excel(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value if math.isfinite(value) else "未提供"
    if isinstance(value, bool):
        return "是" if value else "否"
    text = _text(value, "")
    # Apostrophes force text in Excel, LibreOffice and common CSV conversions.
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        text = "'" + text
    return text[:32767]


def _append(sheet, values):
    sheet.append([_excel(value) for value in values])


def _style(sheet, widths=None):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="16324F")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    from openpyxl.utils import get_column_letter
    for i in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 22)


def _generated(metadata):
    return metadata.get("generated_at") or datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S 北京时间")


def _notes(metadata):
    notes = list(_LIMITS)
    extra = metadata.get("limitations", [])
    if isinstance(extra, str):
        extra = [extra]
    return notes + [_text(item) for item in extra]


def export_rankings(vehicles, events, path, metadata=None):
    """Write a four-sheet review workbook; empty results are valid exports."""
    metadata = dict(metadata or {})
    vehicles = [_mapping(item) for item in (vehicles or [])]
    events = [_mapping(item) for item in (events or [])]
    target = _output_path(path, ".xlsx", metadata)
    book = Workbook()
    ranking = book.active
    ranking.title = "车辆排名"
    _append(ranking, ["排名", "号牌号码", "号牌种类", "车辆类型", "所属单位", "风险等级", "评估状态", "过车记录数", "事件数", "夜间记录比例", "主要行政区域", "首次观测", "末次观测"])
    ordered = sorted(vehicles, key=lambda v: (_level(v.get("risk_level")) or 9, -_number(v.get("event_count")), _text(v.get("plate"), "")))
    for index, vehicle in enumerate(ordered, 1):
        ratio = vehicle.get("night_ratio")
        ratio_text = "未评估" if ratio is None else "{:.1%}".format(_number(ratio))
        _append(ranking, [index, vehicle.get("plate"), vehicle.get("plate_type"), vehicle.get("vehicle_type"), vehicle.get("owner"), _level_text(vehicle), _status(vehicle), vehicle.get("record_count", 0), vehicle.get("event_count", 0), ratio_text, vehicle.get("main_areas"), vehicle.get("first_seen"), vehicle.get("last_seen")])
    if not vehicles:
        _append(ranking, ["", "暂无车辆数据", "", "", "", "未评估", "无可用数据，不能据此判断是否存在异常"])
    _style(ranking, {7: 65, 11: 32})
    evidence = book.create_sheet("事件证据")
    _append(evidence, ["事件编号", "号牌号码", "号牌种类", "事件类型", "事件标题", "风险等级", "规则模型", "开始时间", "结束时间", "规则分值", "触发依据", "关联记录编号", "观测地点", "数据质量与限制"])
    by_identity = {vehicle.get("vehicle_key"): vehicle for vehicle in vehicles if vehicle.get("vehicle_key")}
    plate_types = {}
    for vehicle in vehicles:
        plate_types.setdefault(vehicle.get("plate"), set()).add(vehicle.get("plate_type") or "")
    for event in sorted(events, key=lambda e: (_level(e.get("level")) or 9, _text(e.get("start"), ""), _text(e.get("event_id"), ""))):
        plate_type = event.get("plate_type") or by_identity.get(event.get("vehicle_key"), {}).get("plate_type")
        if not plate_type and len(plate_types.get(event.get("plate"), [])) == 1:
            plate_type = next(iter(plate_types[event.get("plate")]))
        _append(evidence, [event.get("event_id"), event.get("plate"), plate_type, event.get("event_type"), event.get("title"), _LEVEL_NAMES.get(_level(event.get("level")), "未评估"), event.get("model"), event.get("start"), event.get("end"), event.get("score"), event.get("evidence"), event.get("record_keys"), event.get("locations") or event.get("checkpoint") or "请按关联记录编号核对单车报告和原始轨迹", event.get("data_quality") or "仍受卡口覆盖、规则适用条件与源数据质量限制"])
    if not events:
        _append(evidence, ["暂无事件", "", "", "有可用数据时表示已运行规则未命中；无数据或缺必要字段时可能未评估"])
    _style(evidence, {5: 34, 11: 68, 12: 52, 13: 52, 14: 60})
    statistics = book.create_sheet("统计")
    _append(statistics, ["统计项", "数值／说明"])
    levels = Counter(_level(v.get("risk_level")) for v in vehicles)
    for label, value in [("车辆总数", len(vehicles)), ("过车记录数（车辆汇总）", sum(_number(v.get("record_count")) for v in vehicles)), ("疑似事件总数", len(events)), ("一级线索车辆", levels[1]), ("二级线索车辆", levels[2]), ("三级线索车辆", levels[3]), ("未命中规则车辆（仍受数据限制）", sum(1 for v in vehicles if _level(v.get("risk_level")) == 0 and _number(v.get("record_count")) > 0)), ("未评估车辆（无记录或无分析结果）", sum(1 for v in vehicles if _number(v.get("record_count")) <= 0 or _level(v.get("risk_level")) is None))]:
        _append(statistics, [label, value])
    for name, count in sorted(Counter(_text(e.get("event_type")) for e in events).items()):
        _append(statistics, ["事件类型：" + name, count])
    _append(statistics, ["口径", "车辆等级按当前事件和规则结果汇总；事件可能共享观测记录，不应简单累加解释为独立行为次数。"])
    _style(statistics, {1: 40, 2: 100})
    instructions = book.create_sheet("数据说明")
    _append(instructions, ["项目", "内容"])
    _append(instructions, ["生成时间", _generated(metadata)])
    _append(instructions, ["规则版本", metadata.get("rule_version") or RULE_VERSION])
    _append(instructions, ["分析范围", metadata.get("assessment_scope", "仅覆盖本次导出时应用中已有的过车观测及已运行规则")])
    _append(instructions, ["用途", _CAUTION])
    _append(instructions, ["文本安全", "可能被表格软件解释为公式的外部文本已加单引号前缀，按原始文本复核。"])
    for index, note in enumerate(_notes(metadata), 1):
        _append(instructions, ["数据限制 {}".format(index), note])
    _style(instructions, {1: 28, 2: 110})
    book.save(str(target))
    return str(target)


def _paragraph(document, value):
    return document.add_paragraph(_text(value))


def _table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = _text(value)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = _text(value)
    if not rows:
        table.add_row().cells[0].text = "暂无数据（未评估，不能推断不存在异常）"
    return table


def _location(record):
    parts = [_text(record.get(key), "") for key in ("checkpoint", "area", "location_tag")]
    if record.get("longitude") is not None and record.get("latitude") is not None:
        parts.append("坐标 {}，{}".format(record["longitude"], record["latitude"]))
    return "／".join(part for part in parts if part) or "未提供地点"


def export_vehicle_report(detail, path, metadata=None):
    """Write a vehicle report including observed records, event refs and limits."""
    metadata = dict(metadata or {})
    detail = _mapping(detail)
    detail_limits = detail.get("limitations") or []
    if isinstance(detail_limits, str):
        detail_limits = [detail_limits]
    metadata_limits = metadata.get("limitations") or []
    if isinstance(metadata_limits, str):
        metadata_limits = [metadata_limits]
    metadata["limitations"] = list(dict.fromkeys(list(metadata_limits) + list(detail_limits)))
    target = _output_path(path, ".docx", metadata)
    vehicle = _mapping(detail.get("vehicle"))
    records = [_mapping(item) for item in (detail.get("records") or [])]
    events = [_mapping(item) for item in (detail.get("events") or [])]
    history = [_mapping(item) for item in (detail.get("history") or [])]
    records = sorted(records, key=lambda record: (_text(record.get("time"), ""), _text(record.get("record_key"), "")))
    record_map = {record.get("record_key"): record for record in records}
    document = Document()
    document.sections[0].top_margin = Inches(0.7)
    document.sections[0].bottom_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.add_heading("车辆异常轨迹辅助研判报告", 0)
    _paragraph(document, "车辆：{}　生成时间：{}".format(_text(vehicle.get("plate")), _generated(metadata)))
    _paragraph(document, _CAUTION)
    document.add_heading("一、数据范围与车辆画像", 1)
    rows = [("号牌号码", vehicle.get("plate")), ("号牌种类", vehicle.get("plate_type")), ("车辆类型", vehicle.get("vehicle_type")), ("所属单位", vehicle.get("owner")), ("当前风险等级", _level_text(vehicle)), ("评估状态", _status(vehicle)), ("导出轨迹记录数", len(records)), ("当前事件数", len(events)), ("首次观测", vehicle.get("first_seen") or (records[0].get("time") if records else None)), ("末次观测", vehicle.get("last_seen") or (records[-1].get("time") if records else None)), ("主要行政区域", vehicle.get("main_areas")), ("夜间记录比例", "未评估" if vehicle.get("night_ratio") is None else "{:.1%}".format(_number(vehicle["night_ratio"]))), ("规则版本", detail.get("rule_version") or metadata.get("rule_version") or RULE_VERSION), ("分析范围", metadata.get("assessment_scope", "仅依据本报告列出的过车观测；不包含未导入的轨迹"))]
    _table(document, ["项目", "内容"], rows)
    if vehicle.get("assessments"):
        _table(document, ["评估所需数据", "可用性说明"], list(_mapping(vehicle["assessments"]).items()))
    if vehicle.get("unassessed_rules") or detail.get("unassessed_rules"):
        _paragraph(document, "未评估项目：" + _text(vehicle.get("unassessed_rules") or detail.get("unassessed_rules")))
    document.add_heading("二、观测轨迹明细", 1)
    _paragraph(document, "以下为导入的离散观测点，按时间排列，不代表两点之间的完整行驶路线。记录编号用于事件证据溯源。")
    _table(document, ["记录编号", "过车时间", "卡口／行政区域／标签／坐标", "点位类型"], [(r.get("record_key"), r.get("time"), _location(r), r.get("checkpoint_type")) for r in records])
    document.add_heading("三、疑似事件与触发依据", 1)
    if not events:
        if records:
            _paragraph(document, "当前数据和已运行规则未发现异常事件。缺少适用字段的规则可能未评估，本结论不表示车辆不存在异常。")
        else:
            _paragraph(document, "未评估：没有可用过车记录，不能据此判断是否存在异常事件。")
    for index, event in enumerate(events, 1):
        document.add_heading("{}. {}".format(index, _text(event.get("title", event.get("event_type")))), 2)
        _table(document, ["项目", "内容"], [("事件编号", event.get("event_id")), ("事件类型／模型", "{}／{}".format(_text(event.get("event_type")), _text(event.get("model")))), ("线索等级", _LEVEL_NAMES.get(_level(event.get("level")), "未评估")), ("观测时段", "{} 至 {}".format(_text(event.get("start")), _text(event.get("end")))), ("规则参考分值", event.get("score"))])
        evidence = event.get("evidence") or []
        if isinstance(evidence, str):
            evidence = [evidence]
        _paragraph(document, "触发依据：" + (_text(evidence) if evidence else "未提供，应补充核验原始依据"))
        refs = event.get("record_keys") or []
        if isinstance(refs, str):
            refs = [refs]
        ref_rows = []
        for key in refs:
            record = record_map.get(key)
            if record:
                ref_rows.append((key, record.get("time"), _location(record)))
            else:
                ref_rows.append((key, "本报告中未提供该记录", "请在原始数据及其他关联车辆报告中核对"))
        _table(document, ["关联记录编号", "观测时间", "观测地点"], ref_rows)
        _paragraph(document, "证据限制：" + _text(event.get("data_quality") or "仍受卡口覆盖和源数据质量限制；疑似行为需要人工核查。"))
    document.add_heading("四、分析历史", 1)
    _paragraph(document, "下表记录导入或重新分析的操作快照。相同资料的重复分析不代表车辆新增行程，不作为反复异常的证据。")
    history_rows = []
    for item in history:
        operation = item.get("action") or item.get("kind") or item.get("operation") or item.get("reason") or "分析／导入操作"
        history_rows.append((item.get("created_at") or item.get("create_time") or item.get("time"), operation, item.get("batch_id") or item.get("run_id") or item.get("id"), item.get("rule_version") or "未提供", "等级：{}；事件数：{}；记录数：{}".format(_level_text(item), _text(item.get("event_count")), _text(item.get("record_count")))))
    _table(document, ["操作时间", "操作类型", "批次／快照编号", "规则版本", "操作摘要"], history_rows)
    document.add_heading("五、数据限制与未评估范围", 1)
    for note in _notes(metadata):
        _paragraph(document, note)
    document.add_heading("六、人工核查建议", 1)
    for suggestion in ["核对号牌、号牌种类、车型、所属单位及原始过车图像，排除录入错误和同号牌识别混淆。", "按事件关联记录编号复核时间、卡口、坐标、地点标签和出入口属性，核实卡口时钟与覆盖情况。", "对疑似停留、往返、夜间组合等线索，结合路线、正常行程及业务资料复核；需要时依法开展后续核查。", "记录人工核查结论及依据，区分已核实、误报和待补充材料；本报告不自动形成执法结论。"]:
        _paragraph(document, suggestion)
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "辅助研判材料 · 仅供授权业务人员核查"
    document.save(str(target))
    return str(target)


def create_template(path, metadata=None):
    """Create a blank import sheet and a separate, clearly labelled guide."""
    target = _output_path(path, ".xlsx", dict(metadata or {}))
    book = Workbook()
    data = book.active
    data.title = "过车数据"
    fields = ["号牌号码", "过车时间", "车辆类型", "号牌种类", "所属单位", "卡口名称", "行政区域", "地点标签", "经度", "纬度", "点位类型"]
    _append(data, fields)
    _style(data, {1: 20, 2: 25, 5: 30, 6: 30, 8: 32})
    guide = book.create_sheet("填写说明")
    _append(guide, ["字段／项目", "英文名称", "要求及说明"])
    descriptions = [
        ("号牌号码", "plate", "必需。每行一条真实过车观测；同一文件可以包含多辆车。"),
        ("过车时间", "time", "必需。推荐 YYYY-MM-DD HH:MM:SS，按业务发生地本地时间填写。"),
        ("车辆类型", "vehicle_type", "建议填写，例如轻型厢式货车；缺失时不推定车型。"),
        ("号牌种类", "plate_type", "建议填写，用于识别同号牌不同种类。界面按号牌查询，报告保留此字段。"),
        ("所属单位", "owner", "选填；企业群体规则需要可用且准确的所属单位。"),
        ("卡口名称", "checkpoint", "选填；应填写准确点位标识，不能用行政区域代替卡口。"),
        ("行政区域", "area", "选填；用于活动范围和区域往返线索，不单独作为停车依据。"),
        ("地点标签", "location_tag", "选填；使用经核实的地点标签，具体识别词见规则说明。"),
        ("经度", "longitude", "选填；WGS84 十进制度，经度 -180 至 180；与纬度一起填写。"),
        ("纬度", "latitude", "选填；WGS84 十进制度，纬度 -90 至 90；与经度一起填写。"),
        ("点位类型", "checkpoint_type", "选填；高速入口/高速出口及非高速点位属性须与实际一致，具体识别词见规则说明。"),
        ("使用方法", "", "在“过车数据”页从第二行填入观测，不要删除表头；本文件无虚构示例记录。"),
        ("缺少数据", "", "不适用的选填字段留空；不要编造地点、坐标、所属单位或异常标签。字段不足时部分规则可能未评估。"),
        ("数据用途", "", "只导入有权处理的资料；导出报告包含车辆数据，请按业务要求妥善保管。"),
    ]
    for row in descriptions:
        _append(guide, row)
    _style(guide, {1: 24, 2: 25, 3: 110})
    book.save(str(target))
    return str(target)
