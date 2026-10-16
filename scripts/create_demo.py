"""Generate synthetic examples only. Never use production records for releases."""
from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def create_demo(path):
    headers = ["号牌号码", "过车时间", "车辆类型", "号牌种类", "所有人", "卡口名称", "行政区域", "地点标签", "经度", "纬度", "点位类型"]
    rows = []
    def add(plate, at, checkpoint, area, tag="", owner="演示运输企业（虚构）", kind="", vehicle="轻型厢式货车"):
        rows.append([plate, "2026-10-15 " + at, vehicle, "演示号牌", owner, checkpoint, area, tag, None, None, kind])
    # Three vehicles at a labelled sensitive place, with distinct following points.
    for i in range(1, 4):
        plate = "演示A{:04d}".format(i)
        add(plate, "01:00:00", "演示点位S", "演示甲区", "风险区域")
        add(plate, "01:50:00", "演示点位S", "演示甲区", "风险区域")
        add(plate, "02:20:00", "演示点位B", "演示乙区")
    # Complete same-interchange motorway episode.
    add("演示B0001", "10:00:00", "演示南高速出口", "演示甲区", owner="演示高速企业（虚构）", kind="高速出口")
    add("演示B0001", "10:20:00", "演示地面点", "演示甲区", owner="演示高速企业（虚构）", kind="城市道路")
    add("演示B0001", "11:00:00", "演示地面点", "演示甲区", owner="演示高速企业（虚构）", kind="城市道路")
    add("演示B0001", "11:20:00", "演示南高速入口", "演示甲区", owner="演示高速企业（虚构）", kind="高速入口")
    # Frequent known transitions, without same-checkpoint adjacent observations.
    for at, cp, area in [("08:00:00", "演示甲1", "演示甲区"), ("08:20:00", "演示乙1", "演示乙区"), ("08:40:00", "演示甲2", "演示甲区"), ("09:00:00", "演示乙2", "演示乙区"), ("09:20:00", "演示甲3", "演示甲区")]:
        add("演示C0001", at, cp, area, owner="演示往返企业（虚构）")
    # Daytime normal observations: different checkpoints, no suspicious rule.
    add("演示N0001", "09:00:00", "演示普通1", "演示甲区", owner="演示正常企业（虚构）", vehicle="小型客车")
    add("演示N0001", "09:30:00", "演示普通2", "演示乙区", owner="演示正常企业（虚构）", vehicle="小型客车")
    # Nighttime alone is not a suspicious chain.
    add("演示N0002", "23:00:00", "演示夜路1", "演示甲区", owner="演示正常企业（虚构）")
    add("演示N0002", "23:20:00", "演示夜路2", "演示乙区", owner="演示正常企业（虚构）")
    # Clearly labelled hazardous tanker long-interval scenario.
    add("演示D0001", "12:00:00", "演示敏感罐区", "演示丙区", "敏感区域", "演示危化企业（虚构）", vehicle="危险化学品罐式货车")
    add("演示D0001", "16:00:00", "演示敏感罐区", "演示丙区", "敏感区域", "演示危化企业（虚构）", vehicle="危险化学品罐式货车")
    book = Workbook()
    sheet = book.active
    sheet.title = "过车数据"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163A5F")
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 25
    notes = book.create_sheet("填写说明")
    notes.append(["全部数据为软件功能演示而生成的虚构记录，不代表任何真实车辆或企业。"])
    notes.append(["演示日期固定，首页统计范围为已导入数据，不表示电脑当天数据。"])
    notes.append(["包含正常、夜间单独活动、短停组合、高速停返、企业群体、频繁往返和危化品长停。"])
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(target))
    return target


if __name__ == "__main__":
    print(create_demo(Path(__file__).resolve().parent.parent / "samples" / "demo.xlsx"))
