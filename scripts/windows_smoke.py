#!/usr/bin/env python3
"""Run inside the delivered Windows runtime; uses a temporary synthetic DB."""
from __future__ import annotations

import argparse
import json
import os
import platform
import sys
import tempfile
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--show", action="store_true", help="Display the actual Qt main window for two seconds")
    args = parser.parse_args()
    print("TVRA runtime diagnostic")
    print("Platform:", platform.platform())
    print("Python:", sys.version.replace("\n", " "))
    print("Pointer size:", platform.architecture()[0])
    if not args.show:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide2 import QtCore, QtWidgets
    from tvra.service import Service
    from tvra.gui import MainWindow
    from tvra.reports import create_template, export_rankings

    print("Qt:", QtCore.qVersion())
    with tempfile.TemporaryDirectory(prefix="tvra-smoke-") as temporary:
        folder = Path(temporary)
        service = Service(folder / "data")
        template = folder / "template.xlsx"
        create_template(str(template))
        # A blank template is intentionally not a data batch.  Verify its
        # structure, then exercise the import path with the bundled synthetic
        # demo workbook (which contains no real vehicle or enterprise data).
        from openpyxl import load_workbook
        workbook = load_workbook(str(template), read_only=True)
        try:
            assert workbook.sheetnames == ["过车数据", "填写说明"]
        finally:
            workbook.close()
        result = service.import_files([str(ROOT / "samples" / "demo.xlsx")])
        print("Demo import:", json.dumps(result, ensure_ascii=True))
        summary = service.summary()
        print("Summary:", json.dumps(summary, ensure_ascii=True))
        export_rankings(service.list_vehicles(), service.list_events(), str(folder / "report.xlsx"))
        app = QtWidgets.QApplication.instance() or QtWidgets.QApplication(sys.argv[:1])
        window = MainWindow(service)
        if args.show:
            window.show()
        QtCore.QTimer.singleShot(2000 if args.show else 100, app.quit)
        app.exec_()
        window.close()
        service.close()
    print("PASS: SQLite, XLSX and Qt main window completed. No existing user data was changed.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        print("FAIL: see the traceback above. On Windows 7, first check SP1 and the Universal C Runtime update.")
        sys.exit(1)
