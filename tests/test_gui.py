"""Off-screen smoke coverage for the desktop workflow.

The tests exercise the same slots that the user invokes from the six-page
window.  They deliberately use the synthetic demo workbook and temporary
output paths; no production data, dialogs, or packaged executables are used.
"""
from __future__ import annotations

import time
from pathlib import Path

import pytest
from PySide2.QtCore import QDate
from PySide2.QtGui import QFont, QFontMetrics
from PySide2.QtWidgets import QApplication, QFileDialog
from docx import Document
from openpyxl import load_workbook

from scripts.create_demo import create_demo
from tvra.gui import MainWindow
from tvra.service import Service


@pytest.fixture(scope="session")
def qapp():
    # The CI invocation sets QT_QPA_PLATFORM=offscreen.  Keeping the fixture
    # session-scoped avoids creating more than one QApplication in a process.
    app = QApplication.instance() or QApplication([])
    app.setFont(QFont("Noto Sans CJK SC", 9))
    return app


def wait_for_job(window, qapp, timeout_ms=12_000):
    """Pump Qt and give the worker thread scheduling time until it is idle."""
    deadline = time.monotonic() + timeout_ms / 1000
    while window._job is not None and time.monotonic() < deadline:
        # Sleeping briefly releases the GIL so the worker can finish, then
        # process queued completed/finished signals on the GUI thread.
        time.sleep(0.02)
        qapp.processEvents()
    qapp.processEvents()
    assert window._job is None, "后台任务未在测试超时前完成"


@pytest.fixture
def window(tmp_path, qapp):
    instance = MainWindow(Service(tmp_path / "data"))
    # QMessageBox is intentionally not automated in these tests.  User-facing
    # dialogs are covered by capturing their title/text while the underlying
    # slots and worker callbacks are exercised normally.
    instance._info = lambda *args: None
    instance._error = lambda *args: None
    instance.show()
    qapp.processEvents()
    yield instance
    # Tests normally leave the worker idle.  A failed assertion should still
    # leave a usable QApplication for the rest of the suite.
    if instance._job is not None:
        instance._job.wait(5_000)
        qapp.processEvents()
    instance.close()
    qapp.processEvents()


def test_six_pages_demo_import_filters_vehicle_detail_and_font(window, qapp):
    assert window.pages.count() == 6
    metrics = QFontMetrics(qapp.font())
    assert all(metrics.inFont(char) for char in "重点运输车辆异常轨迹智能研判系统")

    window._import_demo()
    wait_for_job(window, qapp)
    assert window.service.summary()["vehicles"] == 8
    assert window.service.summary()["records"] == 24
    assert window.events_table.rowCount() == len(window._events)

    # Event type and keyword filters are wired to Service.list_events().
    window.event_type.setCurrentIndex(1)
    window._refresh_events()
    assert window.events_table.rowCount() > 0
    selected_type = window.event_type.currentData()
    assert all(window.events_table.item(row, 2).text().startswith(selected_type) for row in range(window.events_table.rowCount()))
    window.event_keyword.setText("演示A0001")
    window._refresh_events()
    assert all("演示A0001" in window.events_table.item(row, 1).text() for row in range(window.events_table.rowCount()))

    # Date validation reports a usable error instead of issuing a query.
    messages = []
    window._info = lambda title, text: messages.append((title, text))
    window.date_from_check.setChecked(True)
    window.date_to_check.setChecked(True)
    window.event_date_from.setDate(QDate(2026, 10, 16))
    window.event_date_to.setDate(QDate(2026, 10, 15))
    window._refresh_events()
    assert messages and messages[-1][0] == "日期范围无效"

    # Vehicle detail exposes trajectory, evidence and analysis history.
    vehicle_key = window._vehicles[0]["vehicle_key"]
    window._show_vehicle(vehicle_key)
    assert window._detail["vehicle"]["vehicle_key"] == vehicle_key
    assert window.trajectory_table.rowCount() == len(window._detail["records"])
    assert window.vehicle_history.rowCount() == len(window._detail["history"])


def test_template_xlsx_docx_exports_and_rule_reanalysis(window, qapp, tmp_path, monkeypatch):
    # Export completion is a normal modal notification in interactive use;
    # suppress it here so the off-screen test remains non-interactive.
    window._info = lambda *args: None
    window._import_demo()
    wait_for_job(window, qapp)

    template = tmp_path / "template.xlsx"
    ranking = tmp_path / "ranking.xlsx"
    report = tmp_path / "vehicle.docx"
    choices = iter([template, ranking, report])
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *args: (str(next(choices)), "")))
    window._export_template()
    wait_for_job(window, qapp)
    window._refresh_reports()
    window._export_rankings()
    wait_for_job(window, qapp)
    window.report_table.selectRow(0)
    window._export_selected_report()
    wait_for_job(window, qapp)

    assert load_workbook(template).sheetnames == ["过车数据", "填写说明"]
    assert load_workbook(ranking).sheetnames == ["车辆排名", "事件证据", "统计", "数据说明"]
    assert len(Document(report).tables) >= 2

    # The settings page writes validated values and asks Service to re-run the
    # rules on the existing records in a worker thread.
    window._settings_widgets["short_stop_min"].setValue(35)
    window._settings_widgets["short_stop_max"].setValue(180)
    window._save_settings()
    wait_for_job(window, qapp)
    assert window.service.get_settings()["short_stop_min"] == 35
    assert window.service.summary()["records"] == 24


def test_import_and_worker_errors_are_reported_and_close_is_safe(window, qapp):
    messages = []
    errors = []
    window._info = lambda title, text: messages.append((title, text))
    window._error = lambda *args: errors.append(args)

    window._import_selected()
    assert messages[-1][0] == "尚未选择文件"

    window._start_job("故意失败的测试任务", lambda: 1 / 0, lambda result: None)
    wait_for_job(window, qapp)
    assert errors and "故意失败的测试任务失败" in errors[-1][0]

    # closeEvent refuses to destroy a window while a worker is writing.  Once
    # it finishes, closing releases the Service cleanly and Qt reports no
    # "QThread: Destroyed while thread is still running" abort.
    window._start_job("短暂测试任务", lambda: time.sleep(0.20), lambda result: None)
    window.close()
    assert window.isVisible()
    wait_for_job(window, qapp)
    window.close()
    qapp.processEvents()
    assert not window.isVisible()
