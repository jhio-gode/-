from __future__ import annotations

from openpyxl import Workbook
import pytest

from tvra.service import Service


def workbook(path, rows):
    book = Workbook()
    sheet = book.active
    sheet.append(["号牌号码", "过车时间", "车辆类型", "号牌种类", "所有人", "卡口名称", "行政区域", "地点标签"])
    for row in rows:
        sheet.append(row)
    book.save(str(path))
    return path


def normal_rows(plate="演示Z0001", plate_type="演示"):
    return [[plate, "2026-10-15 09:00:00", "轻型厢式货车", plate_type, "虚构单位", "A", "甲", ""],
            [plate, "2026-10-15 09:30:00", "轻型厢式货车", plate_type, "虚构单位", "B", "乙", ""],
            [plate, "2026-10-15 10:00:00", "轻型厢式货车", plate_type, "虚构单位", "C", "甲", ""]]


def test_normal_cross_region_not_risk_and_repeat_is_idempotent(tmp_path):
    service = Service(tmp_path / "data")
    source = workbook(tmp_path / "input.xlsx", normal_rows())
    first = service.import_files([source])
    assert first["new_records"] == 3
    assert service.list_vehicles()[0]["risk_level"] == 0
    detail = service.vehicle_detail(service.list_vehicles()[0]["vehicle_key"])
    assert len(detail["history"]) == 1
    second = service.import_files([source])
    assert second["new_records"] == 0
    assert service.summary()["records"] == 3
    assert len(service.imports()) == 1
    assert len(service.vehicle_detail(detail["vehicle"]["vehicle_key"])["history"]) == 1
    service.reanalyze()
    assert len(service.vehicle_detail(detail["vehicle"]["vehicle_key"])["history"]) == 1


def test_multicar_and_plate_types_have_independent_identity(tmp_path):
    service = Service(tmp_path / "data")
    rows = normal_rows(plate_type="蓝牌") + normal_rows(plate_type="黄牌") + normal_rows("演示Z0002")
    service.import_files([workbook(tmp_path / "fleet.xlsx", rows)])
    assert service.summary()["vehicles"] == 3
    assert all(v["record_count"] == 3 for v in service.list_vehicles())
    with pytest.raises(ValueError, match="多个号牌种类"):
        service.vehicle_detail("演示Z0001")
    assert service.vehicle_detail("演示Z0001 [蓝牌]")["vehicle"]["plate_type"] == "蓝牌"


def test_cross_batch_ten_second_boundary(tmp_path):
    service = Service(tmp_path / "data")
    row = normal_rows()[0]
    service.import_files([workbook(tmp_path / "one.xlsx", [row])])
    near = list(row)
    near[1] = "2026-10-15 09:00:09"
    exact_boundary = list(row)
    exact_boundary[1] = "2026-10-15 09:00:10"
    result = service.import_files([workbook(tmp_path / "two.xlsx", [near, exact_boundary])])
    assert result["new_records"] == 1
    assert service.summary()["records"] == 2


def test_partial_bad_input_is_reported_not_silently_lost(tmp_path):
    service = Service(tmp_path / "data")
    bad = list(normal_rows()[0])
    bad[1] = "not a timestamp"
    result = service.import_files([workbook(tmp_path / "partial.xlsx", normal_rows() + [bad])])
    assert result["new_records"] == 3
    assert result["error_count"] == 1
    assert service.imports()[0]["error_count"] == 1


def test_enriching_tags_reanalyzes_without_duplicate_passages(tmp_path):
    service = Service(tmp_path / "data")
    first = normal_rows()[0]
    second = list(first)
    second[1] = "2026-10-15 10:00:00"
    source = workbook(tmp_path / "enrich.xlsx", [first, second])
    service.import_files([source])
    assert not any(e["event_type"] == "E04" for e in service.list_events())
    first[-1] = second[-1] = "敏感区域"
    workbook(source, [first, second])
    result = service.import_files([source])
    assert result["new_records"] == 0
    assert result["updated_records"] == 2
    assert service.summary()["records"] == 2
    assert any(e["event_type"] == "E04" for e in service.list_events())


def test_invalid_settings_preserve_configuration_and_data(tmp_path):
    service = Service(tmp_path / "data")
    before = service.get_settings()
    with pytest.raises(ValueError):
        service.save_settings({"short_stop_min": 200, "short_stop_max": 20})
    assert service.get_settings() == before
    with pytest.raises(ValueError):
        service.save_settings({"night_start": True})


def test_failure_is_transactional_and_existing_data_survives(tmp_path, monkeypatch):
    service = Service(tmp_path / "data")
    service.import_files([workbook(tmp_path / "base.xlsx", normal_rows())])
    import tvra.service as module
    def broken(*args, **kwargs):
        raise RuntimeError("simulated analysis failure")
    monkeypatch.setattr(module, "analyze", broken)
    with pytest.raises(RuntimeError):
        service.import_files([workbook(tmp_path / "other.xlsx", normal_rows("演示Z0009"))])
    assert Service(tmp_path / "data").summary()["records"] == 3
    assert len(service.imports()) == 1


def test_export_has_traceable_events_and_no_duplicate_history(tmp_path):
    from scripts.create_demo import create_demo
    from tvra.reports import export_rankings, export_vehicle_report
    from openpyxl import load_workbook
    from docx import Document
    service = Service(tmp_path / "data")
    source = create_demo(tmp_path / "demo.xlsx")
    result = service.import_files([source])
    assert result["error_count"] == 0
    assert service.summary()["vehicles"] == 8
    assert set(e["event_type"] for e in service.list_events()) == {"E01", "E02", "E03", "E04", "E05", "E06"}
    assert any(v["risk_level"] == 0 for v in service.list_vehicles())
    report = export_rankings(service.list_vehicles(), service.list_events(), tmp_path / "rank.xlsx", service.summary())
    book = load_workbook(report)
    assert len(book.sheetnames) == 4
    book.close()
    vehicle = next(v for v in service.list_vehicles() if v["risk_level"] == 1)
    detail = service.vehicle_detail(vehicle["vehicle_key"])
    keys = {r["record_key"] for r in detail["records"]}
    assert all(set(e["record_keys"]) <= keys for e in detail["events"])
    report = export_vehicle_report(detail, tmp_path / "vehicle.docx", service.summary())
    doc = Document(report)
    assert len(doc.tables) >= 2
