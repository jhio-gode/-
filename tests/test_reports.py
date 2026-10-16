from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from tvra.domain import Event, Passage
from tvra.reports import create_template, export_rankings, export_vehicle_report


def document_text(path):
    doc = Document(path)
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )


def test_empty_exports_state_unassessed_without_asserting_no_abnormality(tmp_path):
    ranking = tmp_path / "empty.xlsx"
    report = tmp_path / "empty.docx"
    assert export_rankings([], [], ranking) == str(ranking)
    assert export_vehicle_report({}, report) == str(report)
    workbook = load_workbook(ranking)
    assert workbook.sheetnames == ["车辆排名", "事件证据", "统计", "数据说明"]
    assert workbook["车辆排名"]["F2"].value == "未评估"
    assert "无可用数据" in workbook["车辆排名"]["G2"].value
    assert workbook["统计"]["B2"].value == 0
    text = document_text(report)
    assert "未评估：没有可用过车记录" in text
    assert "不能据此判断是否存在异常事件" in text


def test_xlsx_guards_formula_strings_and_handles_control_characters(tmp_path):
    path = tmp_path / "safe.xlsx"
    vehicles = [{"plate": "=HYPERLINK(\"https://example.test\",\"click\")", "owner": " \t+SUM(1,2)", "vehicle_type": "货车\x00", "record_count": 2, "event_count": 1, "risk_level": 1}]
    events = [{"event_id": "@external", "plate": "-2+3", "event_type": "E01", "title": "=1+1", "level": 1, "score": 80, "evidence": ["=SUM(2,3)"], "record_keys": ["r1"]}]
    export_rankings(vehicles, events, path, {"assessment_scope": "=DDE()"})
    workbook = load_workbook(path, data_only=False)
    assert workbook["车辆排名"]["B2"].value.startswith("'=")
    assert workbook["车辆排名"]["E2"].value.startswith("'")
    assert workbook["车辆排名"]["D2"].value == "货车"
    assert workbook["事件证据"]["J2"].value == 80
    for sheet in workbook:
        assert all(cell.data_type != "f" for row in sheet for cell in row)


def test_docx_contains_record_refs_locations_history_and_limits(tmp_path):
    detail = {
        "vehicle": {"plate": "TEST001", "plate_type": "黄牌", "vehicle_type": "厢式货车", "record_count": 2, "event_count": 1, "risk_level": 2, "night_ratio": 1},
        "records": [Passage("r2", "TEST001", "2026-10-01 23:45:00", checkpoint="合成卡口甲", area="演示区"), Passage("r1", "TEST001", "2026-10-01 23:00:00", checkpoint="合成卡口甲", area="演示区")],
        "events": [Event("e1", "TEST001", "E01", "疑似短停", 2, "2026-10-01 23:00:00", "2026-10-01 23:45:00", 45, ["相邻观测位置接近"], ["r1", "r2", "other-vehicle-ref"], data_quality=["需核对点位坐标"])],
        "history": [{"created_at": "2026-10-02 09:00:00", "action": "重新分析", "batch_id": "demo-batch", "event_count": 1, "record_count": 2}],
        "rule_version": "test-rule-v1",
    }
    target = tmp_path / "vehicle.docx"
    export_vehicle_report(detail, target)
    text = document_text(target)
    for expected in ("test-rule-v1", "黄牌", "分析历史", "demo-batch", "r1", "r2", "other-vehicle-ref", "合成卡口甲", "演示区", "2026-10-01 23:45:00", "相邻观测位置接近", "需核对点位坐标", "本报告中未提供该记录", "重复分析不代表车辆新增行程", "辅助研判材料", "不等于违法事实"):
        assert expected in text


def test_no_event_with_observations_differs_from_no_data(tmp_path):
    target = tmp_path / "observations.docx"
    detail = {"vehicle": {"plate": "TEST002", "record_count": 1, "risk_level": 0, "event_count": 0}, "records": [Passage("r1", "TEST002", "2026-10-01 12:00:00")], "events": []}
    export_vehicle_report(detail, target)
    text = document_text(target)
    assert "当前数据和已运行规则未发现异常事件" in text
    assert "缺少适用字段的规则可能未评估" in text
    assert "未评估：没有可用过车记录" not in text


def test_existing_and_source_files_are_preserved(tmp_path):
    target = tmp_path / "source.xlsx"
    original = b"do not replace source content"
    target.write_bytes(original)
    with pytest.raises(FileExistsError):
        create_template(target)
    with pytest.raises(ValueError, match="源文件"):
        export_rankings([], [], target, {"allow_overwrite": True, "source_paths": [target]})
    with pytest.raises(ValueError, match="源文件"):
        export_rankings([], [], target, {"allow_overwrite": True, "input_paths": str(target)})
    assert target.read_bytes() == original
    # Explicitly confirmed overwrite is supported for an output, never a source.
    export_rankings([], [], target, {"allow_overwrite": True})
    assert load_workbook(target).sheetnames[0] == "车辆排名"


def test_template_contains_only_headers_and_readable_field_guide(tmp_path):
    target = tmp_path / "template.xlsx"
    create_template(target)
    workbook = load_workbook(target)
    assert workbook.sheetnames == ["过车数据", "填写说明"]
    assert workbook["过车数据"].max_row == 1
    assert [c.value for c in workbook["过车数据"][1]][:2] == ["号牌号码", "过车时间"]
    guide_text = " ".join(str(c.value or "") for row in workbook["填写说明"] for c in row)
    assert "plate_type" in guide_text and "checkpoint_type" in guide_text
    assert "不能用行政区域代替卡口" in guide_text


def test_template_round_trip_preserves_all_supported_columns(tmp_path):
    from tvra.importer import load_records

    target = tmp_path / "template.xlsx"
    create_template(target)
    workbook = load_workbook(target)
    workbook["过车数据"].append(["TEST003", "2026-10-01 23:00:00", "厢式货车", "黄牌", "演示单位", "合成卡口", "演示区域", "风险地点", 112.5, 37.8, "高速出口"])
    workbook.save(target)
    imported = load_records([target])
    assert not imported.errors
    assert len(imported.records) == 1
    record = imported.records[0]
    assert record.owner == "演示单位"
    assert record.checkpoint_type == "高速出口"
    assert record.plate_type == "黄牌"
    assert record.longitude == 112.5 and record.latitude == 37.8


def test_exports_reject_wrong_extensions_without_writes(tmp_path):
    target = tmp_path / "wrong.csv"
    with pytest.raises(ValueError):
        export_rankings([], [], target)
    with pytest.raises(ValueError):
        export_vehicle_report({}, target)
    assert not target.exists()
